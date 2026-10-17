#!/usr/bin/env python3
"""Build an editable 16:9 .pptx from an exec deck (HTML or JSON).

Usage:
  python3 build_pptx.py deck.html -o versions/deck-v1.pptx
  python3 build_pptx.py deck.json -o versions/deck-v1.pptx
  python3 build_pptx.py deck.html --dump deck.json   # write the parsed JSON only

Input HTML follows the authoring contract in references/deck-template.html:
one <section class="slide" data-layout="..."> per slide, text marked with
data-field, repeated items wrapped in data-group (data-pick = recommended).
Theme colors and fonts come from the first :root block (hex values only).

JSON shape: {"theme": {"color-bg": "#fff", ...},
             "slides": [{"layout": "big-number", "fields": {"title": ["..."]},
                         "groups": [{"pick": false, "fields": {...}}]}]}

Only dependency: python-pptx. Every text box stays native and editable, and
bullet lists are real PowerPoint bullets (a:buChar with a hanging indent), so
wrapped lines align under the text and PowerPoint's list tools work on them.

Comparison cards may carry data-field="for" and data-field="against" lines
(one of each per option, up to 3 options): the one stated exception to the
3-lines-per-slide cap, so options keep their trade-offs. See card_caps().

Layout warnings (per-slide bullet cap, bullets that will wrap, a big number too
wide for its label) go to stderr. --strict turns any warning into exit code 1,
so a deck cannot be called done while one is open.
"""
import argparse
import json
import os
import re
import sys
from html.parser import HTMLParser

VOID = {"br", "img", "hr", "meta", "link", "input", "source", "wbr"}
DEFAULT_THEME = {
    "color-bg": "#ffffff", "color-surface": "#f3f4f6", "color-rule": "#d9dce1",
    "color-text": "#16181d", "color-text-soft": "#3d424b", "color-text-muted": "#6b717c",
    "color-accent": "#2856c9", "color-accent-strong": "#1d3f96", "color-on-accent": "#ffffff",
    "color-good": "#1f7a4d", "color-risk": "#b4501b",
    "font-display": "Arial", "font-body": "Arial",
}

MAX_POINTS = 3  # per slide, counting every list item on it (cards included)
# The one exception: an options slide keeps each option's trade-off. With up to
# MAX_OPTIONS_WITH_TRADEOFFS options, each may carry one "for" and one "against"
# line (so at most 6 lines); with more options, one line each.
MAX_OPTIONS_WITH_TRADEOFFS = 3
LINE_FIELDS = ("for", "against", "point")  # list-item fields
# Bullet glyph and color per line kind. Shape differs as well as color, so the
# for/against reading survives grayscale print and color blindness.
MARKS = {"point": ("\u2022", "color-accent"), "for": ("+", "color-good"),
         "against": ("\u2212", "color-risk")}

# Point sizes match the HTML template at 1280 px wide (1cqw = 12.8 px = 9.6 pt),
# so a bullet that sits on one line in the HTML render also fits in the pptx.
SIZE = {"point": 18, "col-point": 15, "head": 19, "number": 134, "label": 21, "context": 16}

# Approximate advance widths in em for a bold or regular sans (Arial metrics,
# rounded up). Used to size boxes from the text itself instead of a fixed x.
_EM = {**{c: 0.556 for c in "0123456789$#"}, "%": 0.889, "+": 0.584, "-": 0.333,
       ".": 0.278, ",": 0.278, ":": 0.333, " ": 0.278, "/": 0.278, "(": 0.333,
       ")": 0.333, "x": 0.556, "k": 0.556, "M": 0.833, "W": 0.944, "m": 0.889,
       "w": 0.778, "i": 0.278, "l": 0.278, "j": 0.278, "f": 0.333, "t": 0.333,
       "r": 0.389, "I": 0.278}
WIDTH_SAFETY = 1.08  # estimate only: brand faces run a little wider than Arial
MEASURED_SAFETY = 1.02  # measured from the font file: renderers differ only slightly
MISSING_FONT_SAFETY = 1.2  # a substituted fallback face can be much wider

FONT_DIRS = ["/System/Library/Fonts", "/Library/Fonts", "~/Library/Fonts",
             "/usr/share/fonts", "/usr/local/share/fonts", "~/.fonts",
             "~/.local/share/fonts", "C:/Windows/Fonts"]
FONT_EXT = (".ttf", ".otf", ".ttc", ".otc")
STYLE_WORDS = {"regular", "roman", "book", "normal", "medium", "bold", "italic", "oblique",
               "light", "thin", "black", "heavy", "semibold", "demibold", "extrabold",
               "ultrabold", "condensed", "narrow", "mt", "ps"}
REGULAR_STYLES = {"regular", "roman", "book", "normal", ""}
_font_files = None
_face_cache = {}  # path -> [(path, index, family, style)]
_font_cache = {}


def _key(name):
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _faces_from_filename(path):
    """Fallback when no font library is available: "Arial Bold.ttf" -> ("Arial", "Bold").

    Style words are stripped from the end of the file name; whatever is left is
    the family. "ArialHB.ttc" stays family "ArialHB", so it can never pass for Arial.
    """
    words = re.findall(r"[A-Za-z0-9]+", os.path.splitext(os.path.basename(path))[0])
    style = []
    while len(words) > 1 and words[-1].lower() in STYLE_WORDS:
        style.insert(0, words.pop())
    return [(path, 0, " ".join(words), " ".join(style) or "Regular")]


def _read_faces(path):
    """Every face in a font file as (path, index, family, style), from the name table.

    .ttc/.otc collections hold several faces; each gets its own index. Uses
    fontTools when installed (it reads the typographic family, name IDs 16/17),
    else Pillow (ships with python-pptx), else the file name.
    """
    if path in _face_cache:
        return _face_cache[path]
    faces = []
    try:
        from fontTools.ttLib import TTCollection, TTFont
        fonts = (TTCollection(path, lazy=True).fonts
                 if path.lower().endswith((".ttc", ".otc")) else [TTFont(path, lazy=True)])
        for idx, ft in enumerate(fonts):
            nm = ft["name"]
            fam = nm.getDebugName(16) or nm.getDebugName(1)
            sty = nm.getDebugName(17) or nm.getDebugName(2) or ""
            if fam:
                faces.append((path, idx, fam, sty))
    except ImportError:
        try:
            from PIL import ImageFont
            for idx in range(64):
                try:
                    fam, sty = ImageFont.truetype(path, 12, index=idx).getname()
                except (OSError, ValueError):
                    break
                faces.append((path, idx, fam or "", sty or ""))
        except ImportError:
            faces = _faces_from_filename(path)
    except Exception:  # unreadable or odd font file: skip it, never crash the build
        faces = []
    _face_cache[path] = faces
    return faces


def pick_face(name, bold, faces):
    """Choose (path, index) for family `name` from (path, index, family, style) candidates.

    The family must match exactly (case and spacing ignored): "Arial" never
    matches "Arial Hebrew", "Arial Unicode MS" or "Arial Black". Within the
    family, an exact style wins (Bold, or Regular/Roman/Book), then the closest
    upright style. Ties go to a single-face file over a collection, then to the
    path in sort order, so the choice is stable and never depends on path length.
    """
    key = _key(name)
    fam = [f for f in faces if _key(f[2]) == key]
    if not fam:
        return None

    def rank(f):
        st = _key(f[3])
        slanted = "italic" in st or "oblique" in st
        if bold:
            exact = st == "bold"
            near = "bold" in st and not slanted and not any(w in st for w in ("semi", "demi", "extra", "ultra"))
            loose = "bold" in st and not slanted
        else:
            exact = st in REGULAR_STYLES
            near = st == "medium"
            loose = not slanted and "bold" not in st
        tier = 0 if exact else 1 if near else 2 if loose else 3
        return (tier, f[0].lower().endswith((".ttc", ".otc")), f[0], f[1])

    best = min(fam, key=rank)
    return best[0], best[1]


def find_font(name, bold=False):
    """Return (path, index) of the installed face for family `name`, or None.

    Files whose name starts like the family are read first (cheap); if none of
    them is that exact family, every installed font is read, because a file
    name need not resemble its family name.
    """
    global _font_files
    ck = (name, bold)
    if ck in _font_cache:
        return _font_cache[ck]
    if _font_files is None:
        _font_files = []
        for d in FONT_DIRS:
            for root, _, files in os.walk(os.path.expanduser(d)):
                _font_files += sorted(os.path.join(root, f) for f in files
                                      if f.lower().endswith(FONT_EXT))
    head = _key(name)[:4]
    likely = [f for f in _font_files if _key(os.path.basename(f)).startswith(head)]
    hit = pick_face(name, bold, [face for f in likely for face in _read_faces(f)])
    if hit is None:
        seen = set(likely)
        rest = [f for f in _font_files if f not in seen]
        hit = pick_face(name, bold, [face for f in rest for face in _read_faces(f)])
    _font_cache[ck] = hit
    return hit


def resolve_font(stack):
    """First installed face in a CSS font stack; else the first name given."""
    names = [n.strip().strip("\"'") for n in stack.split(",") if n.strip()]
    names = [n for n in names if n not in ("sans-serif", "serif", "system-ui", "monospace")]
    for n in names:
        if find_font(n):
            return n
    return names[0] if names else "Arial"


def text_width_in(text, size_pt, bold=False, font=None):
    """Rendered width of one line of text, in inches.

    Measured from the font file when `font` is installed here (Pillow ships
    with python-pptx); otherwise estimated from Arial-like metrics, padded
    more when the font is missing because the viewer will substitute one.
    """
    if font:
        found = find_font(font, bold)
        if found:
            try:
                from PIL import ImageFont
                face = ImageFont.truetype(found[0], 1000, index=found[1])
                return face.getlength(text) / 1000 * size_pt / 72 * MEASURED_SAFETY
            except (ImportError, OSError):
                pass
    em = 0.0
    for ch in text:
        if ch in _EM:
            em += _EM[ch]
        elif ch.isupper():
            em += 0.722
        else:
            em += 0.556
    if bold:
        em *= 1.06
    pad = MISSING_FONT_SAFETY if font else WIDTH_SAFETY
    return em * size_pt / 72 * pad


def card_caps(sd, limit=MAX_POINTS):
    """(slide_total, [per-card cap]) line budget for one slide.

    Every slide gets `limit` lines in total, cards included. The one exception
    is a comparison with up to MAX_OPTIONS_WITH_TRADEOFFS options: a card that
    states its cost in an "against" line may keep it next to its one "for" (or
    plain point) line, so the cap there is 2 per option. A decision slide that
    drops the cost of each option no longer shows the trade-off being decided.
    Extra plain points never earn the exception.
    """
    groups = sd.get("groups", [])
    if sd.get("layout") == "comparison" and groups:
        tradeoffs = len(groups) <= MAX_OPTIONS_WITH_TRADEOFFS
        caps = [2 if tradeoffs and g.get("fields", {}).get("against") else 1 for g in groups]
        return max(limit, sum(caps)), caps
    return limit, [limit] * len(groups)


def group_lines(g):
    """A card's list lines as [(field, text)] in keep order: its lead line (first
    "for", else first point), its first "against", then everything else."""
    f = g.get("fields", {})
    rest = [("for", x) for x in f.get("for", [])] + [("point", x) for x in f.get("point", [])]
    lines = rest[:1] + [("against", x) for x in f.get("against", [])[:1]]
    return lines + rest[1:] + [("against", x) for x in f.get("against", [])[1:]]


def enforce_point_cap(deck, limit=MAX_POINTS):
    """Trim each slide to its line budget (see card_caps). Returns warnings; mutates the deck.

    Slide-level points are kept first, then card lines are handed out one per
    card per round (round-robin), so every option keeps its lead line before
    any option keeps its "against" line.
    """
    warnings = []
    for i, sd in enumerate(deck.get("slides", []), 1):
        f = sd.get("fields", {})
        groups = [g.setdefault("fields", {}) for g in sd.get("groups", [])]
        lines = [group_lines({"fields": g}) for g in groups]
        total_cap, caps = card_caps(sd, limit)
        total = len(f.get("point", [])) + sum(len(ls) for ls in lines)
        if total <= total_cap and all(len(ls) <= c for ls, c in zip(lines, caps)):
            continue
        budget = total_cap
        if f.get("point"):
            f["point"] = f["point"][:budget]
            budget -= len(f["point"])
        kept = [[] for _ in groups]
        rnd = 0
        while budget > 0 and any(len(ls) > rnd and c > rnd for ls, c in zip(lines, caps)):
            for gi, (ls, c) in enumerate(zip(lines, caps)):
                if budget and len(ls) > rnd and c > rnd:
                    kept[gi].append(ls[rnd])
                    budget -= 1
            rnd += 1
        for g, k in zip(groups, kept):
            for field in LINE_FIELDS:
                g.pop(field, None)
            for field, text in k:
                g.setdefault(field, []).append(text)
        if sd.get("layout") == "comparison" and groups:
            rule = (f"one lead line per option, plus one \"against\" line when there are "
                    f"{MAX_OPTIONS_WITH_TRADEOFFS} options or fewer")
        else:
            rule = f"{limit} per slide (cards included)"
        warnings.append(f"slide {i}: {total} bullets, the cap is {rule}; extras dropped. "
                        f"Cut them in the HTML.")
    return warnings


class DeckParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.slides, self.stack, self.css = [], [], []
        self.slide = self.group = None
        self.in_style = False

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "style":
            self.in_style = True
        if tag in VOID:
            if tag == "br":
                self._text(" ")
            return
        frame = {"tag": tag, "slide": False, "group": False, "field": None, "buf": []}
        if tag == "section" and "slide" in (a.get("class") or "").split():
            self.slide = {"layout": a.get("data-layout", "bottom-line"), "fields": {}, "groups": []}
            frame["slide"] = True
        elif self.slide is not None and "data-group" in a:
            self.group = {"pick": "data-pick" in a, "fields": {}}
            self.slide["groups"].append(self.group)
            frame["group"] = True
        if self.slide is not None and a.get("data-field"):
            frame["field"] = a["data-field"]
        self.stack.append(frame)

    def handle_endtag(self, tag):
        if tag == "style":
            self.in_style = False
        if tag in VOID:
            return
        while self.stack:
            frame = self.stack.pop()
            if frame["field"]:
                text = re.sub(r"\s+", " ", "".join(frame["buf"])).strip()
                target = self.group if self.group is not None else self.slide
                if text and target is not None:
                    target["fields"].setdefault(frame["field"], []).append(text)
            if frame["group"]:
                self.group = None
            if frame["slide"]:
                self.slides.append(self.slide)
                self.slide = None
            if frame["tag"] == tag:
                break

    def handle_data(self, data):
        if self.in_style:
            self.css.append(data)
        else:
            self._text(data)

    def _text(self, data):
        for frame in self.stack:
            if frame["field"]:
                frame["buf"].append(data)


def parse_html(src):
    p = DeckParser()
    p.feed(src)
    theme = dict(DEFAULT_THEME)
    m = re.search(r":root\s*{([^}]*)}", "".join(p.css))
    if m:
        for name, val in re.findall(r"--([\w-]+)\s*:\s*([^;]+);", m.group(1)):
            val = val.strip()
            if name.startswith("font-"):
                theme[name] = resolve_font(val)
            elif re.fullmatch(r"#[0-9a-fA-F]{6}", val):
                theme[name] = val
    return {"theme": theme, "slides": p.slides}


# ----------------------------------------------------------------------------- pptx
def build(deck, out):
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt
        from lxml.etree import SubElement
    except ImportError:
        sys.exit("python-pptx is not installed: pip install python-pptx")

    t = dict(DEFAULT_THEME, **deck.get("theme", {}))
    rgb = lambda k: RGBColor.from_string(t[k].lstrip("#").upper())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    L, W = 0.8, 13.333 - 1.6  # left margin and content width, inches
    warnings = []
    slide_no = [0]

    def bulletize(para, size, char="\u2022", color="color-accent"):
        """Real list formatting: bullet glyph in a:buChar, hanging indent via marL/indent."""
        pPr = para._p.get_or_add_pPr()
        hang = int(Pt(size) * 1.1)
        pPr.set("marL", str(hang))
        pPr.set("indent", str(-hang))
        for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buNone", "a:buChar", "a:buAutoNum"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        clr = SubElement(pPr, qn("a:buClr"))
        SubElement(clr, qn("a:srgbClr"), val=t[color].lstrip("#").upper())
        SubElement(pPr, qn("a:buFont"), typeface="Arial")
        SubElement(pPr, qn("a:buChar"), char=char)
        return hang / 914400  # inches

    for k in ("font-display", "font-body"):
        if not find_font(t[k]):
            warnings.append(f"font \"{t[k]}\" is not installed here, so renders substitute a "
                            f"different face; pass --font with an installed one")

    def box(s, x, y, w, h, text, size, color="color-text", bold=False, font="font-body",
            bullet=False, one_line=False, wrap=True, marks=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.space_after = Pt(size * 0.6)
            room = w - 0.2  # default left + right insets
            if bullet:
                mark = marks[i] if marks else MARKS["point"]
                room -= bulletize(para, size, *mark)
            if (bullet or one_line) and text_width_in(line, size, bold, t[font]) > room:
                warnings.append(f"slide {slide_no[0]}: line will wrap: \"{line}\" "
                                f"(shorten it to one line)")
            r = para.add_run()
            r.text = line
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, t[font]
            r.font.color.rgb = rgb(color)
        return tb

    def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        sh.line.fill.background()
        sh.shadow.inherit = False
        style = sh._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:  # drop theme effects (shadows) some viewers still apply
            sh._element.remove(style)
        return sh

    def first(d, k, default=""):
        return (d.get(k) or [default])[0]

    def bullets(d):
        return list(d.get("point", []))

    def card_lines(d):
        lines = group_lines({"fields": d})
        return [x for _, x in lines], [MARKS[k] for k, _ in lines]

    warnings.extend(enforce_point_cap(deck))
    for n, sd in enumerate(deck["slides"], 1):
        slide_no[0] = n
        s = prs.slides.add_slide(blank)
        f, groups, lay = sd.get("fields", {}), sd.get("groups", []), sd.get("layout")
        bg = "color-accent-strong" if lay == "title" else "color-bg"
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        ink = "color-on-accent" if lay == "title" else "color-text"

        if lay == "title":
            box(s, L, 2.0, W, 0.4, first(f, "kicker").upper(), 13, ink, True)
            box(s, L, 2.5, W * 0.8, 2.0, first(f, "title"), 48, ink, True, "font-display")
            box(s, L, 4.6, W * 0.8, 0.8, first(f, "subtitle"), 20, ink)
            box(s, L, 5.6, W, 0.4, first(f, "meta"), 13, ink)
        else:
            box(s, L, 0.55, W, 0.4, first(f, "kicker").upper(), 12, "color-accent", True)
            title_size = 34 if lay == "bottom-line" else 30
            box(s, L, 0.95, W * 0.85, 1.3, first(f, "title"), title_size, "color-text", True, "font-display")

        if lay == "bottom-line":
            rect(s, L, 2.6, 0.08, max(1.0, 0.3 + 0.45 * len(bullets(f))), "color-accent")  # fit the list
            box(s, L + 0.35, 2.7, W - 0.5, 2.6, bullets(f), SIZE["point"], "color-text-soft", bullet=True)
        elif lay == "big-number":
            num, label = first(f, "number"), first(f, "label")
            nx, nw, lx, lw, size = big_number_layout(num, label, L, W, font=t["font-display"])
            if size < SIZE["number"]:
                warnings.append(f"slide {n}: number \"{num}\" shrunk to {size} pt to leave "
                                f"room for its label; consider a shorter number")
            box(s, nx, 2.3, nw, 2.6, num, size, "color-accent", True, "font-display", wrap=False)
            box(s, lx, 3.45, lw, 1.3, label, SIZE["label"], "color-text", True)
            box(s, L, 5.1, W, 0.8, first(f, "context"), SIZE["context"], "color-text-soft")
        elif lay == "comparison":
            gw = (W - 0.4 * (len(groups) - 1)) / max(len(groups), 1)
            most = max([len(group_lines(g)) for g in groups] or [0])
            badge = 0.5 if any(g.get("pick") for g in groups) else 0
            card_h = min(3.7, max(1.8, 0.25 + badge + 0.8 + 0.45 * most + 0.35))  # fit content
            for i, g in enumerate(groups):
                x = L + i * (gw + 0.4)
                rect(s, x, 2.4, gw, card_h, "color-surface")
                rect(s, x, 2.4, gw, 0.07, "color-accent" if g.get("pick") else "color-rule")
                y = 2.65
                if g.get("pick"):
                    pill = rect(s, x + 0.3, y, 1.7, 0.35, "color-accent-strong", MSO_SHAPE.ROUNDED_RECTANGLE)
                    pill.text_frame.text = "RECOMMENDED"
                    run = pill.text_frame.paragraphs[0].runs[0]
                    run.font.size, run.font.bold, run.font.name = Pt(10), True, t["font-body"]
                    run.font.color.rgb = rgb("color-on-accent")
                    y += 0.5
                box(s, x + 0.3, y, gw - 0.6, 0.6, first(g["fields"], "head"), SIZE["head"], "color-text",
                    True, "font-display", one_line=True)
                texts, marks = card_lines(g["fields"])
                box(s, x + 0.3, y + 0.8, gw - 0.6, 2.2, texts, SIZE["col-point"],
                    "color-text-soft", bullet=True, marks=marks)
        elif lay == "timeline":
            rect(s, L, 3.05, W, 0.04, "color-rule")
            gw = W / max(len(groups), 1)
            for i, g in enumerate(groups):
                x = L + i * gw
                dot = rect(s, x, 2.88, 0.38, 0.38, "color-accent" if g.get("pick") else "color-bg", MSO_SHAPE.OVAL)
                dot.line.color.rgb = rgb("color-accent")
                dot.line.width = Pt(3)
                box(s, x, 3.5, gw - 0.2, 0.4, first(g["fields"], "when").upper(), 14, "color-accent", True)
                box(s, x, 3.95, gw - 0.2, 1.4, first(g["fields"], "what"), 19, "color-text", True)
        elif lay == "decision":
            ask = rect(s, L, 2.5, W, 1.0, "color-accent-strong", MSO_SHAPE.ROUNDED_RECTANGLE)
            ask.adjustments[0] = 0.12
            tf = ask.text_frame
            tf.word_wrap = True
            tf.text = first(f, "ask")
            run = tf.paragraphs[0].runs[0]
            run.font.size, run.font.bold, run.font.name = Pt(24), True, t["font-display"]
            run.font.color.rgb = rgb("color-on-accent")
            box(s, L, 3.9, W, 2.0, bullets(f), SIZE["point"], "color-text-soft", bullet=True)

        if lay != "title":
            rect(s, L, 6.7, W, 0.01, "color-rule")
            box(s, L, 6.75, W - 1.0, 0.4, first(f, "source"), 10, "color-text-muted")
            box(s, L + W - 0.6, 6.75, 0.6, 0.4, str(n), 10, "color-text-muted")
        if f.get("notes"):
            s.notes_slide.notes_text_frame.text = " ".join(f["notes"])

    prs.save(out)
    print(f"wrote {out}: {len(deck['slides'])} slides")
    return warnings


def big_number_layout(num, label, left, width, size=None, gap=0.3, label_min=3.2, font=None):
    """Place the label from the number's estimated width, not a fixed x.

    Returns (number_x, number_w, label_x, label_w, number_size_pt). If the
    number would leave less than label_min inches for the label, its size
    steps down until it fits.
    """
    size = size or SIZE["number"]
    while True:
        nw = text_width_in(num, size, bold=True, font=font) + 0.2  # + text box insets
        if width - nw - gap >= label_min or size <= 60:
            break
        size -= 6
    lx = left + nw + gap
    return left, nw, lx, left + width - lx, size


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("src", help="deck .html or .json")
    ap.add_argument("-o", "--out", help="output .pptx path (put it in versions/)")
    ap.add_argument("--font", help="font for every text box (one installed where the deck is opened)")
    ap.add_argument("--dump", help="write the parsed deck JSON to this path")
    ap.add_argument("--strict", action="store_true",
                    help="exit 1 if any layout warning is raised (use before calling a deck done)")
    args = ap.parse_args()
    with open(args.src, encoding="utf-8") as fh:
        raw = fh.read()
    deck = json.loads(raw) if args.src.endswith(".json") else parse_html(raw)
    if not deck.get("slides"):
        sys.exit("no slides found: check the section.slide markup")
    warnings = enforce_point_cap(deck)
    if args.font:
        deck.setdefault("theme", {}).update({"font-display": args.font, "font-body": args.font})
    if args.dump:
        with open(args.dump, "w", encoding="utf-8") as fh:
            json.dump(deck, fh, indent=2)
        print(f"wrote {args.dump}")
    if not (args.out or args.dump):
        ap.error("give -o and/or --dump")
    if args.out:
        warnings += build(deck, args.out)
    for w in warnings:
        print(f"warning: {w}", file=sys.stderr)
    if warnings and args.strict:
        sys.exit(f"{len(warnings)} layout warning(s); fix the HTML and rebuild")


if __name__ == "__main__":
    main()
