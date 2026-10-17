"""Tests for build_pptx.py.

Run from this folder:
  PYTHONDONTWRITEBYTECODE=1 uv run --with python-pptx --with pytest \
    pytest -q -p no:cacheprovider test_build_pptx.py
"""
import copy
import os
import subprocess
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_pptx as bp  # noqa: E402

TEMPLATE = os.path.join(HERE, "..", "references", "deck-template.html")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def three_option_deck(points_per_option=2):
    groups = [
        {"pick": name == "B", "fields": {
            "head": [f"{name}: option {name.lower()}"],
            "point": [f"{name} point {k}" for k in range(1, points_per_option + 1)]}}
        for name in "ABC"
    ]
    return {"theme": {}, "slides": [
        {"layout": "comparison", "fields": {"title": ["Pick B"], "source": ["Source: test"]},
         "groups": groups},
        {"layout": "big-number", "fields": {"title": ["Volume grew"], "number": ["+38%"],
                                             "label": ["tickets quarter over quarter, same 9 agents"],
                                             "source": ["Source: test"]}},
    ]}


def build(deck, tmp_path):
    from pptx import Presentation
    out = str(tmp_path / "deck.pptx")
    warnings = bp.build(copy.deepcopy(deck), out)
    return Presentation(out), warnings


def bullet_paragraphs(slide):
    paras = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                pPr = p._p.pPr
                if pPr is not None and pPr.find(A + "buChar") is not None:
                    paras.append(p)
    return paras


# ---------------------------------------------------------------- native bullets
def test_bullets_are_native_with_hanging_indent(tmp_path):
    deck = bp.parse_html(open(TEMPLATE, encoding="utf-8").read())
    prs, _ = build(deck, tmp_path)
    bottom_line = prs.slides[1]
    paras = bullet_paragraphs(bottom_line)
    assert len(paras) == 3
    for p in paras:
        pPr = p._p.pPr
        assert int(pPr.get("marL")) > 0
        assert int(pPr.get("indent")) == -int(pPr.get("marL"))
        assert pPr.find(A + "buChar").get("char") == "•"
        assert not p.text.startswith("•"), "bullet must not be typed into the text"


def test_no_typed_bullet_characters_anywhere(tmp_path):
    deck = bp.parse_html(open(TEMPLATE, encoding="utf-8").read())
    prs, _ = build(deck, tmp_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert "•" not in shape.text_frame.text


# ---------------------------------------------------------------- per-slide cap
def test_cap_counts_all_cards_on_the_slide(tmp_path):
    prs, warnings = build(three_option_deck(points_per_option=2), tmp_path)
    paras = bullet_paragraphs(prs.slides[0])
    assert len(paras) == 3, "6 card bullets must be cut to 3 for the whole slide"
    # round-robin: every option keeps its first point
    assert sorted(p.text for p in paras) == ["A point 1", "B point 1", "C point 1"]
    assert any("slide 1" in w and "cap" in w for w in warnings)


def test_cap_keeps_slide_points_first():
    deck = {"slides": [{"layout": "bottom-line",
                        "fields": {"point": ["a", "b", "c", "d"]}, "groups": []}]}
    warnings = bp.enforce_point_cap(deck)
    assert deck["slides"][0]["fields"]["point"] == ["a", "b", "c"]
    assert len(warnings) == 1


def test_template_obeys_cap_and_has_three_option_comparison():
    deck = bp.parse_html(open(TEMPLATE, encoding="utf-8").read())
    assert bp.enforce_point_cap(copy.deepcopy(deck)) == []
    comp = [sd for sd in deck["slides"] if sd["layout"] == "comparison"][0]
    assert len(comp["groups"]) == 3
    css = open(TEMPLATE, encoding="utf-8").read()
    assert "grid-auto-flow: column" in css and "grid-template-columns: 1fr 1fr" not in css


def test_template_builds_without_warnings(tmp_path):
    deck = bp.parse_html(open(TEMPLATE, encoding="utf-8").read())
    _, warnings = build(deck, tmp_path)
    # a missing font is a machine fact, not a template defect
    assert [w for w in warnings if not w.startswith("font ")] == []


# ---------------------------------------------------------------- one-line bullets
def test_long_card_bullet_warns_it_will_wrap(tmp_path):
    deck = three_option_deck(points_per_option=1)
    deck["slides"][0]["groups"][0]["fields"]["point"] = [
        "Billing response time goes back to eleven hours on average"]
    _, warnings = build(deck, tmp_path)
    assert any("will wrap" in w for w in warnings)


def test_long_option_head_warns_it_will_wrap(tmp_path):
    deck = three_option_deck(points_per_option=1)
    deck["slides"][0]["groups"][2]["fields"]["head"] = ["C: roll out to every queue now"]
    _, warnings = build(deck, tmp_path)
    assert any("will wrap" in w and "every queue" in w for w in warnings)


def test_number_box_never_wraps(tmp_path):
    prs, _ = build(three_option_deck(), tmp_path)
    nums = [sh for sh in prs.slides[1].shapes if sh.has_text_frame and sh.text_frame.text == "+38%"]
    assert nums and nums[0].text_frame.word_wrap is False


def test_missing_font_is_padded_and_warned(tmp_path):
    plain = bp.text_width_in("+38%", 134, bold=True)
    missing = bp.text_width_in("+38%", 134, bold=True, font="NoSuchFace Zz")
    assert missing > plain
    deck = three_option_deck(points_per_option=1)
    deck["theme"] = {"font-display": "NoSuchFace Zz"}
    _, warnings = build(deck, tmp_path)
    assert any(w.startswith("font ") and "NoSuchFace" in w for w in warnings)


def test_resolve_font_skips_uninstalled_faces():
    assert bp.resolve_font('"NoSuchFace Zz", sans-serif') == "NoSuchFace Zz"
    installed = bp.resolve_font('"NoSuchFace Zz", Arial, sans-serif')
    assert installed in ("Arial", "NoSuchFace Zz")  # Arial when this machine has it


# ---------------------------------------------------------------- big number
@pytest.mark.parametrize("num", ["+38%", "2.5h", "18", "$1.2M"])
def test_label_starts_after_measured_number(num):
    left, width = 0.8, 11.733
    nx, nw, lx, lw, size = bp.big_number_layout(num, "label", left, width)
    est = bp.text_width_in(num, size, bold=True)
    assert lx >= nx + est, "label must not overlap the number"
    assert lx - (nx + est) < 0.8, "label must not float far from the number"
    assert lw >= 3.2


def test_wide_and_narrow_numbers_get_different_label_x():
    _, _, lx_wide, _, _ = bp.big_number_layout("+38%", "x", 0.8, 11.733)
    _, _, lx_narrow, _, _ = bp.big_number_layout("18", "x", 0.8, 11.733)
    assert lx_wide - lx_narrow > 1.0


def test_very_long_number_shrinks_to_leave_label_room():
    _, _, _, lw, size = bp.big_number_layout("$12,400,000", "x", 0.8, 11.733)
    assert size < bp.SIZE["number"]
    assert lw >= 3.2


# ---------------------------------------------------------------- trade-off lines
def tradeoff_deck(n_options=3, extra_point=False):
    groups = []
    for name in "ABCD"[:n_options]:
        fields = {"head": [f"{name}: option"], "for": [f"{name} plus"], "against": [f"{name} cost"]}
        if extra_point:
            fields["point"] = [f"{name} extra"]
        groups.append({"pick": name == "B", "fields": fields})
    return {"theme": {}, "slides": [{"layout": "comparison",
                                     "fields": {"title": ["Pick B"], "source": ["Source: test"]},
                                     "groups": groups}]}


def test_three_options_keep_for_and_against_without_warning(tmp_path):
    deck = tradeoff_deck(3)
    assert bp.enforce_point_cap(copy.deepcopy(deck)) == []
    prs, warnings = build(deck, tmp_path)
    paras = bullet_paragraphs(prs.slides[0])
    assert sorted(p.text for p in paras) == sorted(
        f"{n} {k}" for n in "ABC" for k in ("plus", "cost"))
    assert not [w for w in warnings if "cap" in w]


def test_for_and_against_get_distinct_native_marks(tmp_path):
    prs, _ = build(tradeoff_deck(3), tmp_path)
    marks = {p.text.split()[1]: p._p.pPr.find(A + "buChar").get("char")
             for p in bullet_paragraphs(prs.slides[0])}
    assert marks["plus"] == "+" and marks["cost"] == "\u2212"


def test_tradeoff_exception_does_not_admit_extra_points():
    deck = tradeoff_deck(3, extra_point=True)
    warnings = bp.enforce_point_cap(deck)
    assert len(warnings) == 1
    for g in deck["slides"][0]["groups"]:
        f = g["fields"]
        assert f.get("for") and f.get("against") and "point" not in f


def test_four_options_get_one_line_each():
    deck = tradeoff_deck(4)
    warnings = bp.enforce_point_cap(deck)
    assert warnings
    for g in deck["slides"][0]["groups"]:
        f = g["fields"]
        assert f.get("for") and "against" not in f


def test_plain_points_on_options_still_capped_at_three():
    deck = three_option_deck(points_per_option=2)
    bp.enforce_point_cap(deck)
    assert sum(len(g["fields"].get("point", [])) for g in deck["slides"][0]["groups"]) == 3


def test_template_options_carry_tradeoffs():
    deck = bp.parse_html(open(TEMPLATE, encoding="utf-8").read())
    comp = [sd for sd in deck["slides"] if sd["layout"] == "comparison"][0]
    for g in comp["groups"]:
        assert len(g["fields"].get("for", [])) == 1
        assert len(g["fields"].get("against", [])) == 1


# ---------------------------------------------------------------- font matching
FAKE_FACES = [
    ("/System/Library/Fonts/ArialHB.ttc", 0, "Arial Hebrew", "Regular"),
    ("/System/Library/Fonts/ArialHB.ttc", 1, "Arial Hebrew", "Bold"),
    ("/System/Library/Fonts/Arial Unicode.ttf", 0, "Arial Unicode MS", "Regular"),
    ("/x/Arial Black.ttf", 0, "Arial Black", "Regular"),
    ("/x/Arial Narrow Bold.ttf", 0, "Arial Narrow", "Bold"),
    ("/a/much/longer/path/Supplemental/Arial Bold Italic.ttf", 0, "Arial", "Bold Italic"),
    ("/a/much/longer/path/Supplemental/Arial Bold.ttf", 0, "Arial", "Bold"),
    ("/a/much/longer/path/Supplemental/Arial.ttf", 0, "Arial", "Regular"),
]


def test_pick_face_matches_family_exactly_not_by_prefix():
    assert bp.pick_face("Arial", True, FAKE_FACES)[0].endswith("/Arial Bold.ttf")
    assert bp.pick_face("Arial", False, FAKE_FACES)[0].endswith("/Arial.ttf")
    assert bp.pick_face("arial hebrew", True, FAKE_FACES) == ("/System/Library/Fonts/ArialHB.ttc", 1)
    assert bp.pick_face("Ari", True, FAKE_FACES) is None


def test_pick_face_ignores_path_length_and_order():
    faces = list(reversed(FAKE_FACES))
    assert bp.pick_face("Arial", True, faces)[0].endswith("/Arial Bold.ttf")


def test_pick_face_prefers_single_file_over_collection_on_tie():
    faces = [("/a/Fam.ttc", 2, "Fam", "Bold"), ("/zzz/long/path/Fam-Bold.ttf", 0, "Fam", "Bold")]
    assert bp.pick_face("Fam", True, faces) == ("/zzz/long/path/Fam-Bold.ttf", 0)


def test_pick_face_falls_back_to_nearest_upright_style():
    faces = [("/a/F-Italic.ttf", 0, "F", "Bold Italic"), ("/b/F-Semi.ttf", 0, "F", "SemiBold")]
    assert bp.pick_face("F", True, faces)[0] == "/b/F-Semi.ttf"
    faces = [("/a/F-Italic.ttf", 0, "F", "Italic"), ("/b/F-Medium.ttf", 0, "F", "Medium")]
    assert bp.pick_face("F", False, faces)[0] == "/b/F-Medium.ttf"


def test_filename_fallback_keeps_family_exact():
    assert bp._faces_from_filename("/f/Arial Bold.ttf")[0][2:] == ("Arial", "Bold")
    assert bp._faces_from_filename("/f/ArialHB.ttc")[0][2] == "ArialHB"


@pytest.mark.parametrize("bold,style", [(True, "Bold"), (False, "Regular")])
def test_real_arial_resolves_to_arial_family(bold, style):
    from PIL import ImageFont
    found = bp.find_font("Arial", bold)
    if found is None:
        pytest.skip("Arial is not installed on this machine")
    family, got_style = ImageFont.truetype(found[0], 12, index=found[1]).getname()
    assert family == "Arial", f"{found} is {family}, not Arial (Hebrew/Unicode must not match)"
    assert got_style == style


# ---------------------------------------------------------------- CLI gate
def test_strict_exits_nonzero_on_warnings(tmp_path):
    deck_json = tmp_path / "deck.json"
    import json
    deck_json.write_text(json.dumps(three_option_deck(points_per_option=2)))
    run = subprocess.run([sys.executable, os.path.join(HERE, "build_pptx.py"), str(deck_json),
                          "-o", str(tmp_path / "out.pptx"), "--strict"],
                         capture_output=True, text=True,
                         env=dict(os.environ, PYTHONDONTWRITEBYTECODE="1"))
    assert run.returncode == 1
    assert "cap" in run.stderr
